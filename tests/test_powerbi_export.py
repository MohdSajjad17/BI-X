from pathlib import Path

from openpyxl import load_workbook

from bix.services.powerbi_export import write_xlsx


def _meta():
    table = {
        "name": "Dim Test", "hidden": False,
        "columns": [{"name": "ID", "type": "Data", "dataType": "int64", "formatString": None,
                     "hidden": False, "summarizeBy": None, "sourceColumn": "ID", "sortByColumn": None,
                     "displayFolder": None, "dataCategory": None, "isKey": False, "description": None,
                     "expression": None, "lineageTag": None}],
        "measures": [{"name": "Count", "displayFolder": None, "formatString": "0", "hidden": False,
                      "description": None, "expression": "COUNTROWS('Dim Test')", "lineageTag": None}],
        "partitions": [], "hierarchies": [], "calculationGroup": None, "description": None, "lineageTag": None,
    }
    return {
        "input": "Supervisor.zip",
        "project": {"name": "Supervisor"},
        "semanticModel": {
            "format": "TMDL", "info": {}, "tables": [table],
            "relationships": [{"fromTable": "Dim Test", "fromColumn": "ID", "toTable": "Dim Test", "toColumn": "ID",
                               "fromCardinality": "many", "toCardinality": "one", "crossFilteringBehavior": "oneDirection",
                               "isActive": True, "id": "r1"}],
            "daxQueries": [{"name": "DAX1", "content": "EVALUATE ROW()"}],
        },
        "powerQuery": [{"query": "Dim Test", "kind": "Table", "language": "M", "mode": "Import",
                        "queryGroup": None, "dependsOn": [], "sourceObjects": [], "sourceFunctions": [], "code": "let x=1 in x"}],
        "parameters": [{"name": "P", "currentValue": "1", "type": "Text", "defaultValue": "1",
                        "allowedValues": None, "required": False, "queryGroup": None}],
        "report": {
            "info": {}, "pages": [{"order": 1, "id": "p", "name": "Page", "activeOnOpen": True, "visibility": "Visible",
                                   "pageType": "Standard", "width": 1, "height": 1, "displayOption": "FitToPage",
                                   "visualCount": 1, "groupCount": 0, "pageFilterCount": 0}],
            "visuals": [{"page": "Page", "id": "v", "type": "barChart", "title": "Test", "text": None, "group": None,
                         "hidden": False, "x": 0, "y": 0, "width": 1, "height": 1, "z": 0, "tabOrder": 0,
                         "fieldCount": 1, "fieldsUsed": "Dim Test[ID]", "visualFilterCount": 0}],
            "visualFields": [{"page": "Page", "visualId": "v", "visualType": "barChart", "title": "Test", "role": "Category",
                              "table": "Dim Test", "field": "ID", "fieldKind": "Column", "aggregationCode": None,
                              "displayName": "ID", "modelCheck": "OK"}],
            "filters": [{"level": "Page", "owner": "Page", "filterId": "f", "table": "Dim Test", "field": "ID",
                         "fieldKind": "Column", "filterType": "Basic", "values": "1", "hiddenInViewMode": False,
                         "lockedInViewMode": False, "modelCheck": "OK"}],
            "unresolvedReferences": [],
        },
    }


def test_reference_workbook_shape(tmp_path: Path):
    output = write_xlsx(_meta(), tmp_path / "metadata.xlsx")
    wb = load_workbook(output, read_only=True, data_only=True)
    assert wb.sheetnames == [
        "Summary", "Tables", "Columns", "Measures", "Relationships", "Power Query",
        "Parameters", "Pages", "Visuals", "Visual Fields", "Filters", "Unresolved Refs", "DAX Queries",
    ]
    assert [c.value for c in next(wb["Visuals"].iter_rows(min_row=1, max_row=1))] == [
        "Page", "Visual ID", "Visual Type", "Title", "Text content", "Group", "Hidden",
        "X", "Y", "Width", "Height", "Z", "Tab Order", "# Fields", "Fields used", "# Visual filters",
    ]
