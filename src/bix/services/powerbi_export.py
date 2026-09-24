from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def table_category(t: dict[str, Any]) -> str:
    name = t["name"]
    if t.get("calculationGroup"):
        return "Calculation group"
    if name.lower().startswith("fact"):
        return "Fact"
    if name.lower().startswith("dim"):
        return "Dimension"
    if t.get("measures") and len(t.get("columns", [])) <= 1 and len(t.get("measures", [])) >= 5:
        return "Measure table"
    return "Helper / Other"


def _meta_from_ir(ir: dict[str, Any]) -> dict[str, Any]:
    ref = dict(ir.get("extensions", {}).get("reference_metadata", {}))
    project = dict(ref.get("project", {}))
    project.setdefault("name", ir.get("identity", {}).get("name"))
    return {**ref, "input": ir.get("identity", {}).get("source"), "project": project}


def flat(meta: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    m, r = meta.get("semanticModel"), meta.get("report")
    out: dict[str, list[dict[str, Any]]] = {}
    if m:
        out["Tables"] = [{
            "Table": t["name"], "Category": table_category(t), "Hidden": t["hidden"],
            "# Columns": len(t["columns"]), "# Measures": len(t["measures"]),
            "Partition Type": ", ".join(sorted({p["type"] for p in t["partitions"] if p.get("type")})) or None,
            "Mode": ", ".join(sorted({p["mode"] for p in t["partitions"] if p.get("mode")})) or None,
            "Query Group": ", ".join(sorted({p["queryGroup"] for p in t["partitions"] if p.get("queryGroup")})) or None,
            "Description": t.get("description"), "Lineage Tag": t.get("lineageTag"),
        } for t in m["tables"]]
        out["Columns"] = [{
            "Table": t["name"], "Column": c["name"], "Type": c.get("type"), "Data Type": c.get("dataType"),
            "Format String": c.get("formatString"), "Hidden": c.get("hidden"), "Summarize By": c.get("summarizeBy"),
            "Source Column": c.get("sourceColumn"), "Sort By Column": c.get("sortByColumn"),
            "Display Folder": c.get("displayFolder"), "Data Category": c.get("dataCategory"),
            "Is Key": c.get("isKey"), "Description": c.get("description"),
            "DAX Expression (calculated)": c.get("expression"), "Lineage Tag": c.get("lineageTag"),
        } for t in m["tables"] for c in t["columns"]]
        out["Measures"] = [{
            "Table": t["name"], "Measure": x["name"], "Display Folder": x.get("displayFolder"),
            "Format String": x.get("formatString"), "Hidden": x.get("hidden"), "Description": x.get("description"),
            "DAX Expression": x.get("expression"), "Lineage Tag": x.get("lineageTag"),
        } for t in m["tables"] for x in t["measures"]]
        out["Relationships"] = [{
            "From Table": x["fromTable"], "From Column": x["fromColumn"], "To Table": x["toTable"], "To Column": x["toColumn"],
            "Cardinality": f'{x["fromCardinality"]}-to-{x["toCardinality"]}',
            "Cross-filter Direction": "Both" if x["crossFilteringBehavior"] == "bothDirections" else "Single",
            "Active": x["isActive"], "Relationship ID": x["id"],
        } for x in m["relationships"]]
        out["Power Query"] = [{
            "Query / Table": q["query"], "Kind": q["kind"], "Language": q["language"], "Mode": q["mode"],
            "Query Group": q["queryGroup"], "Depends On (other queries)": ", ".join(q["dependsOn"]),
            "Source objects navigated": q["sourceObjects"], "Source functions": ", ".join(q["sourceFunctions"]), "Code": q["code"],
        } for q in meta.get("powerQuery", [])]
        out["Parameters"] = [{
            "Parameter": p["name"], "Current Value": p["currentValue"], "Type": p["type"], "Default Value": p["defaultValue"],
            "Allowed Values": p["allowedValues"], "Required": p["required"], "Query Group": p["queryGroup"],
        } for p in meta.get("parameters", [])]
        out["DAX Queries"] = [{"Name": d["name"], "Content": d["content"]} for d in m.get("daxQueries", [])]
    if r:
        out["Pages"] = [{
            "Order": p["order"], "Page ID": p["id"], "Page Name": p["name"], "Active on open": p["activeOnOpen"],
            "Visibility": p["visibility"], "Page Type": p["pageType"], "Width": p["width"], "Height": p["height"],
            "Display Option": p["displayOption"], "# Visuals (excl. groups)": p["visualCount"], "# Groups": p["groupCount"],
            "# Page filters": p["pageFilterCount"],
        } for p in r["pages"]]
        out["Visuals"] = [{
            "Page": v["page"], "Visual ID": v["id"], "Visual Type": v["type"], "Title": v["title"], "Text content": v["text"],
            "Group": v["group"], "Hidden": v["hidden"], "X": v["x"], "Y": v["y"], "Width": v["width"], "Height": v["height"],
            "Z": v["z"], "Tab Order": v["tabOrder"], "# Fields": v["fieldCount"], "Fields used": v["fieldsUsed"],
            "# Visual filters": v["visualFilterCount"],
        } for v in r["visuals"]]
        out["Visual Fields"] = [{
            "Page": x["page"], "Visual ID": x["visualId"], "Visual Type": x["visualType"], "Title": x["title"],
            "Role (field well)": x["role"], "Table": x["table"], "Field": x["field"], "Field Kind": x["fieldKind"],
            "Aggregation code": x["aggregationCode"], "Display name in visual": x["displayName"], "Model check": x["modelCheck"],
        } for x in r["visualFields"]]
        out["Filters"] = [{
            "Level": x["level"], "Owner": x["owner"], "Filter ID": x["filterId"], "Table": x["table"], "Field": x["field"],
            "Field Kind": x["fieldKind"], "Filter Type": x["filterType"], "Values / Condition literals": x["values"],
            "Hidden": x.get("hiddenInViewMode"), "Locked": x.get("lockedInViewMode"), "Model check": x["modelCheck"],
        } for x in r["filters"]]
        out["Unresolved Refs"] = [{
            "Table referenced": x["table"], "Field referenced": x["field"], "Kind": x["kind"], "Problem": x["problem"],
            "# References": x["references"], "Pages": "; ".join(x["pages"]),
        } for x in r["unresolvedReferences"]]
    return out


def summary_rows(meta: dict[str, Any]) -> list[tuple[str, Any]]:
    m, r = meta.get("semanticModel"), meta.get("report")
    rows = [("Tool", "pbip_reverse_engineer.py v1.0"), ("Input", meta.get("input")), ("Project name", meta["project"]["name"])]
    if m:
        i = m["info"]
        rows += [
            ("MODEL", ""), ("Format", m["format"]), ("Display name", i.get("displayName")), ("Logical ID", i.get("logicalId")),
            ("Compatibility level", i.get("compatibilityLevel")), ("Culture", i.get("culture")),
            ("Data source version", i.get("defaultPowerBIDataSourceVersion")), ("Tables", len(m["tables"])),
            ("Columns", sum(len(t["columns"]) for t in m["tables"])),
            ("  calculated columns", sum(1 for t in m["tables"] for c in t["columns"] if c["type"] == "Calculated")),
            ("Measures", sum(len(t["measures"]) for t in m["tables"])), ("Relationships", len(m["relationships"])),
            ("Parameters", len(meta.get("parameters", []))), ("Power Query queries", len(meta.get("powerQuery", []))),
            ("Storage modes", ", ".join(sorted({p["mode"] for t in m["tables"] for p in t["partitions"] if p.get("mode")})) or None),
        ]
    if r:
        i = r["info"]
        rows += [
            ("REPORT", ""), ("Format", i.get("pbirFormat")), ("Display name", i.get("displayName")), ("Logical ID", i.get("logicalId")),
            ("Base theme", i.get("baseTheme")), ("Custom theme", i.get("customTheme")), ("Pages", len(r["pages"])),
            ("Visual containers", len(r["visuals"])), ("Field references", len(r["visualFields"])),
            ("  unresolved vs model", len([x for x in r["visualFields"] if x["modelCheck"] not in ("OK", "Model not available")])),
            ("Filters", len(r["filters"])), ("DAX Queries", len(m.get("daxQueries", [])) if m else 0),
        ]
    return rows


def write_xlsx(meta: dict[str, Any], path: Path) -> Path:
    F = flat(meta)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    alt = PatternFill("solid", fgColor="F2F5FA")
    bad = PatternFill("solid", fgColor="FDE9E7")
    line = Border(bottom=Side(style="thin", color="D0D7E2"))
    illegal = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

    def put(sheet, row, col, value, **font):
        if isinstance(value, bool):
            value = "Yes" if value else "No"
        if isinstance(value, (list, dict)):
            value = json.dumps(value, ensure_ascii=False)
        if isinstance(value, str):
            value = illegal.sub("", value)[:32000]
        cell = sheet.cell(row=row, column=col, value=value)
        if isinstance(value, str) and value[:1] in "=+-@":
            cell.data_type = "s"
        font.setdefault("size", 10)
        cell.font = Font(name="Arial", **font)
        return cell

    put(ws, 1, 1, f'Power BI Project Metadata — {meta["project"]["name"]}', bold=True, size=14, color="1F3864")
    for row, (key, value) in enumerate(summary_rows(meta), 3):
        if value == "":
            for col in (1, 2):
                ws.cell(row=row, column=col).fill = hdr_fill
            put(ws, row, 1, key, bold=True, color="FFFFFF").fill = hdr_fill
        else:
            put(ws, row, 1, key, bold=True)
            put(ws, row, 2, value).alignment = Alignment(horizontal="left", wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 90
    ws.sheet_view.showGridLines = False

    order = ["Tables", "Columns", "Measures", "Relationships", "Power Query", "Parameters", "Pages", "Visuals", "Visual Fields", "Filters", "Unresolved Refs", "DAX Queries"]
    code_cols = {"Code", "DAX Expression", "Content"}
    for name in order:
        rows = F.get(name, [])
        sheet = wb.create_sheet(name)
        headers = list(rows[0].keys()) if rows else ["(none)"]
        for col, header in enumerate(headers, 1):
            cell = put(sheet, 1, col, header, bold=True, color="FFFFFF")
            cell.fill = hdr_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row, data in enumerate(rows, 2):
            flagged = name == "Unresolved Refs" or data.get("Model check") not in (None, "OK", "Model not available")
            for col, header in enumerate(headers, 1):
                cell = put(sheet, row, col, data.get(header))
                cell.alignment = Alignment(vertical="top", wrap_text=header in code_cols)
                cell.border = line
                if flagged:
                    cell.fill = bad
                elif row % 2 == 1:
                    cell.fill = alt
        for col, header in enumerate(headers, 1):
            longest = max([len(str(header))] + [len(str(x.get(header))) for x in rows[:300] if x.get(header) is not None])
            sheet.column_dimensions[get_column_letter(col)].width = min(max(10, longest + 2), 70 if header in code_cols else 48)
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        sheet.row_dimensions[1].height = 24
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
