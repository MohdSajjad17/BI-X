#!/usr/bin/env python3
"""Extract structured and lossless metadata from a Tableau .twb workbook.

Outputs:
  tableau_metadata_full.json  - complete recursive TWB XML representation
  tableau_metadata_full.xlsx  - flattened metadata sheets for BI-X analysis
  manifest.json               - extraction counts and output locations

Usage:
  python tableau_metadata_extractor.py workbook.twb [output_dir]
"""
from __future__ import annotations

import collections
import json
import sys
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_EXCEL_CELL_LENGTH = 32000
DEFAULT_OUTPUT = "tableau_metadata_extracted"


def tag(el):
    return el.tag.split("}", 1)[-1]


def txt(el):
    return (el.text or "").strip() if el is not None else ""


def attr(el, *names):
    if el is None:
        return None
    for name in names:
        if name in el.attrib:
            return el.attrib[name]
    return None


def child(el, name):
    if el is None:
        return None
    return next((x for x in el if tag(x) == name), None)


def descendants(el, name):
    return [x for x in el.iter() if tag(x) == name]


def nodes(root, name):
    return descendants(root, name)


def xml_dict(el):
    out = {"tag": tag(el), "attributes": dict(el.attrib)}
    value = txt(el)
    if value:
        out["text"] = value
    if len(el):
        out["children"] = [xml_dict(x) for x in el]
    return out


def dedupe(rows):
    seen, out = set(), []
    for row in rows:
        key = json.dumps(row, sort_keys=True, ensure_ascii=False, default=str)
        if key not in seen:
            seen.add(key)
            out.append(row)
    return out


def extract(source: Path):
    tree = ET.parse(source)
    root = tree.getroot()

    datasources, fields, calculations, parameters, filters = [], [], [], [], []
    groups, sets, bins, aliases, hierarchies = [], [], [], [], []
    connections, relations = [], []
    worksheets, views, dashboards, zones, stories, actions = [], [], [], [], [], []
    column_instances, metadata_records, formats, style_rules, cards = [], [], [], [], []

    for ds in nodes(root, "datasource"):
        name = attr(ds, "name")
        datasources.append({
            "Datasource": name,
            "Caption": attr(ds, "caption"),
            "Has Extract": bool(nodes(ds, "extract")),
            "Connection Count": len(nodes(ds, "connection")),
            "Column Count": len(nodes(ds, "column")),
            "Parameter Count": len(nodes(ds, "param")),
        })
        for c in nodes(ds, "column"):
            calc = child(c, "calculation")
            formula = attr(calc, "formula") if calc is not None else None
            if formula is None and calc is not None:
                formula = txt(calc)
            fields.append({
                "Datasource": name, "Field": attr(c, "name"),
                "Caption": attr(c, "caption"),
                "Datatype": attr(c, "datatype", "data-type"),
                "Role": attr(c, "role"), "Type": attr(c, "type"),
                "Semantic Role": attr(c, "semantic-role"),
                "Default Aggregation": attr(c, "default-aggregation"),
                "Calculation": formula, "Ordinal": attr(c, "ordinal"),
                "Hidden": attr(c, "hidden"),
                "Param Domain": attr(c, "param-domain-type"),
            })
            if calc is not None:
                calculations.append({
                    "Datasource": name, "Field": attr(c, "name"),
                    "Caption": attr(c, "caption"),
                    "Calculation Type": attr(calc, "class"),
                    "Formula": formula,
                    **{f"calc.{k}": v for k, v in calc.attrib.items() if k != "formula"},
                })
        for con in nodes(ds, "connection"):
            connections.append({"Datasource": name, **con.attrib})
        for rel in nodes(ds, "relation"):
            relations.append({"Datasource": name, **rel.attrib, "Relation SQL": txt(rel)})
        for p in nodes(ds, "param"):
            parameters.append({"Datasource": name, **p.attrib, "Text": txt(p)})
        for f in nodes(ds, "filter"):
            filters.append({"Scope": "Datasource", "Datasource": name, **f.attrib, "Text": txt(f)})

    for p in nodes(root, "param"):
        parameters.append({"Datasource": None, **p.attrib, "Text": txt(p)})
    for f in nodes(root, "filter"):
        filters.append({"Scope": "Workbook/XML", **f.attrib, "Text": txt(f)})
    for e, target in (("group", groups), ("set", sets), ("bin", bins), ("alias", aliases), ("hierarchy", hierarchies)):
        for x in nodes(root, e):
            target.append({"Tag": e, **x.attrib, "Text": txt(x)})

    for ws in nodes(root, "worksheet"):
        name = attr(ws, "name")
        worksheets.append({
            "Worksheet": name,
            "Column Count": len(nodes(ws, "column")),
            "Column Instance Count": len(nodes(ws, "column-instance")),
            "Filter Count": len(nodes(ws, "filter")),
            "Mark Count": len(nodes(ws, "mark")),
            "Run Count": len(nodes(ws, "run")),
            "Card Count": len(nodes(ws, "card")),
            "Format Count": len(nodes(ws, "format")),
            "Zone Count": len(nodes(ws, "zone")),
        })
        for z in nodes(ws, "zone"):
            zones.append({"Container": name, "Container Type": "Worksheet", **z.attrib})
        for v in nodes(ws, "view"):
            views.append({"Worksheet": name, **v.attrib})

    for db in nodes(root, "dashboard"):
        name = attr(db, "name")
        dashboards.append({
            "Dashboard": name,
            "Zone Count": len(nodes(db, "zone")),
            "Action Count": len(nodes(db, "action")),
            "Filter Count": len(nodes(db, "filter")),
            "Size Count": len(nodes(db, "size")),
        })
        for z in nodes(db, "zone"):
            zones.append({"Container": name, "Container Type": "Dashboard", **z.attrib})

    for st in nodes(root, "story"):
        stories.append({"Story": attr(st, "name"), **st.attrib})
    for a in nodes(root, "action"):
        actions.append({"Tag": "action", **a.attrib, "Text": txt(a)})
    for x in nodes(root, "column-instance"):
        column_instances.append({"Tag": "column-instance", **x.attrib, "Text": txt(x)})
    for x in nodes(root, "metadata-record"):
        metadata_records.append({"Tag": "metadata-record", **x.attrib, "Text": txt(x)})
    for x in nodes(root, "format"):
        formats.append({"Tag": "format", **x.attrib, "Text": txt(x)})
    for x in nodes(root, "style-rule"):
        style_rules.append({"Tag": "style-rule", **x.attrib, "Text": txt(x)})
    for x in nodes(root, "card"):
        cards.append({"Tag": "card", **x.attrib, "Text": txt(x)})

    parameters, filters = dedupe(parameters), dedupe(filters)
    counts = collections.Counter(tag(x) for x in root.iter())
    inventory = [{"XML Tag": k, "Count": v} for k, v in counts.most_common()]

    summary = [
        {"Metric": "Source file", "Value": source.name},
        {"Metric": "File size (bytes)", "Value": source.stat().st_size},
        {"Metric": "Workbook XML version", "Value": root.attrib.get("version")},
        {"Metric": "Original version", "Value": root.attrib.get("original-version")},
        {"Metric": "Source build", "Value": root.attrib.get("source-build")},
        {"Metric": "Source platform", "Value": root.attrib.get("source-platform")},
        {"Metric": "XML tag count", "Value": len(counts)},
        {"Metric": "Total XML nodes", "Value": sum(counts.values())},
        {"Metric": "Datasources", "Value": len(datasources)},
        {"Metric": "Worksheets", "Value": len(worksheets)},
        {"Metric": "Dashboards", "Value": len(dashboards)},
        {"Metric": "Stories", "Value": len(stories)},
        {"Metric": "Actions", "Value": len(actions)},
        {"Metric": "Calculated fields", "Value": len(calculations)},
        {"Metric": "Parameters", "Value": len(parameters)},
        {"Metric": "Filters", "Value": len(filters)},
        {"Metric": "Groups", "Value": len(groups)},
        {"Metric": "Sets", "Value": len(sets)},
        {"Metric": "Bins", "Value": len(bins)},
        {"Metric": "Aliases", "Value": len(aliases)},
        {"Metric": "Hierarchies", "Value": len(hierarchies)},
        {"Metric": "Zones", "Value": len(zones)},
        {"Metric": "Column instances", "Value": len(column_instances)},
        {"Metric": "Metadata records", "Value": len(metadata_records)},
    ]

    return {
        "root": root, "summary": summary, "inventory": inventory,
        "datasources": datasources, "connections": connections, "relations": relations,
        "fields": fields, "calculations": calculations, "parameters": parameters,
        "filters": filters, "worksheets": worksheets, "views": views,
        "dashboards": dashboards, "zones": zones, "stories": stories,
        "actions": actions, "groups": groups, "sets": sets, "bins": bins,
        "aliases": aliases, "hierarchies": hierarchies,
        "column_instances": column_instances, "metadata_records": metadata_records,
        "formats": formats, "style_rules": style_rules, "cards": cards,
    }


def write_excel(meta, output):
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    alt_fill = PatternFill("solid", fgColor="F2F5FA")

    def sheet(name, rows):
        ws = wb.create_sheet(name[:31])
        rows = rows or [{"(none)": ""}]
        headers = list(rows[0])
        for j, h in enumerate(headers, 1):
            c = ws.cell(1, j, h)
            c.fill, c.font = header_fill, header_font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, row in enumerate(rows, 2):
            for j, h in enumerate(headers, 1):
                value = row.get(h)
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                if value is not None:
                    value = str(value)
                    if len(value) > MAX_EXCEL_CELL_LENGTH:
                        value = value[:MAX_EXCEL_CELL_LENGTH] + "…[truncated; see JSON]"
                c = ws.cell(i, j, value)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2 == 1:
                    c.fill = alt_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        for j, h in enumerate(headers, 1):
            longest = len(str(h))
            for row in rows[:500]:
                v = row.get(h)
                if v is not None:
                    longest = max(longest, min(60, len(str(v))))
            ws.column_dimensions[get_column_letter(j)].width = min(max(12, longest + 2), 60)

    order = [
        ("Workbook Info", meta["summary"]), ("XML Inventory", meta["inventory"]),
        ("Data Sources", meta["datasources"]), ("Connections", meta["connections"]),
        ("Relations", meta["relations"]), ("Fields", meta["fields"]),
        ("Calculations", meta["calculations"]), ("Parameters", meta["parameters"]),
        ("Filters", meta["filters"]), ("Worksheets", meta["worksheets"]),
        ("Views", meta["views"]), ("Dashboards", meta["dashboards"]),
        ("Zones", meta["zones"]), ("Stories", meta["stories"]),
        ("Actions", meta["actions"]), ("Groups", meta["groups"]),
        ("Sets", meta["sets"]), ("Bins", meta["bins"]), ("Aliases", meta["aliases"]),
        ("Hierarchies", meta["hierarchies"]), ("Column Instances", meta["column_instances"]),
        ("Metadata Records", meta["metadata_records"]), ("Formats", meta["formats"]),
        ("Style Rules", meta["style_rules"]), ("Cards", meta["cards"]),
    ]
    for name, rows in order:
        sheet(name, rows)
    wb.save(output)


def write_outputs(meta, source, outdir):
    outdir.mkdir(parents=True, exist_ok=True)
    json_path = outdir / "tableau_metadata_full.json"
    xlsx_path = outdir / "tableau_metadata_full.xlsx"
    manifest_path = outdir / "manifest.json"
    json_path.write_text(json.dumps(xml_dict(meta["root"]), ensure_ascii=False, indent=2), encoding="utf-8")
    write_excel(meta, xlsx_path)
    manifest = {
        "source": str(source), "source_size_bytes": source.stat().st_size,
        "workbook_attributes": meta["root"].attrib,
        "counts": {k: len(meta[k]) for k in (
            "datasources", "connections", "relations", "fields", "calculations", "parameters",
            "filters", "worksheets", "views", "dashboards", "zones", "stories", "actions",
            "groups", "sets", "bins", "aliases", "hierarchies", "column_instances", "metadata_records",
            "formats", "style_rules", "cards")},
        "outputs": {"json": str(json_path), "excel": str(xlsx_path)},
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return json_path, xlsx_path, manifest_path


def main(argv=None):
    argv = argv or sys.argv[1:]
    if not argv:
        raise SystemExit("Usage: python tableau_metadata_extractor.py workbook.twb [output_dir]")
    source = Path(argv[0]).resolve()
    if not source.exists():
        raise SystemExit(f"File not found: {source}")
    outdir = Path(argv[1]).resolve() if len(argv) > 1 else source.parent / DEFAULT_OUTPUT
    meta = extract(source)
    outputs = write_outputs(meta, source, outdir)
    print("Extraction completed")
    print(f"Source: {source}")
    for row in meta["summary"]:
        print(f'{row["Metric"]}: {row["Value"]}')
    print("Outputs:")
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
